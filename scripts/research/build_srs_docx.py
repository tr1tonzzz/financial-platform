"""Build and validate the SRS Word document from the authoritative Markdown.

Run using the Codex bundled Python runtime. Layout QA uses the packaged
document renderer via render_proposal_windows.py on this Windows workspace.
"""
from __future__ import annotations

import json
import re
from pathlib import Path
from collections import Counter

from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH, WD_TAB_ALIGNMENT, WD_TAB_LEADER
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT
from docx.oxml import OxmlElement
from docx.oxml.ns import qn

ROOT = Path(__file__).resolve().parents[2]
SOURCE = ROOT / 'docs/research-platform/22-srs-dac-ta-yeu-cau-phan-mem.md'
OUTPUT = ROOT / 'docs/research-platform/reports/SRS-FAP-01-v3.1.docx'
QA = ROOT / '_plan_render/srs-v3.1'


def validate(text: str) -> dict:
    declared = re.findall(r'^\*\*((?:FR|BR|DR|IR|NFR)\d{2})\b', text, re.M)
    counts = Counter(re.sub(r'\d+$', '', item) for item in declared)
    assert len(declared) == len(set(declared)), 'Repeated requirement declaration'
    assert counts == {'FR': 38, 'BR': 12, 'DR': 8, 'IR': 5, 'NFR': 10}, counts
    tests = set(re.findall(r'^\| (TC\d{2}) —', text, re.M))
    uses = set(re.findall(r'^### [\d.]+ (UC\d{2}) —', text, re.M))
    assert len(tests) == 31 and len(uses) == 8, (tests, uses)
    assert set(re.findall(r'\bTC\d{2}\b', text)) == tests
    assert set(re.findall(r'\bUC\d{2}\b', text)) == uses
    matrix = re.findall(r'^\| ((?:FR|DR|IR|NFR)\d{2}) \| (.+)$', text, re.M)
    expected = {item for item in declared if not item.startswith('BR')}
    assert len(matrix) == len(expected) and {row[0] for row in matrix} == expected
    for rid, cols in matrix:
        assert re.search(r'S\d{2}', cols), f'{rid}: missing source'
        assert re.search(r'UC\d{2}', cols), f'{rid}: missing use case'
        assert re.search(r'TC\d{2}', cols), f'{rid}: missing verification'
    assert set(re.findall(r'\bBR\d{2}\b', '\n'.join(c for _, c in matrix))) == {
        item for item in declared if item.startswith('BR')
    }, 'Business rules not allocated'
    for match in re.finditer(r'\[[^\]]+\]\(([^)]+)\)', text):
        target = match.group(1)
        if not target.startswith('https://'):
            assert (SOURCE.parent / target).exists(), target
    assert '\ufffd' not in text
    extensions = re.findall(r'^\*\*(E\d{2})\b', text, re.M)
    assert set(extensions) == {f'E{i:02d}' for i in range(1, 11)} and len(extensions) == 10
    extension_matrix = re.findall(r'^\| (E\d{2}) \| (.+)$', text, re.M)
    assert len(extension_matrix) == 10 and {r for r, _ in extension_matrix} == set(extensions)
    extension_tests = set(re.findall(r'\bP\d{2}\b', text.split('\n## 15.',1)[1]))
    assert extension_tests == {f'P{i:02d}' for i in range(1, 12)}
    for rid, cols in extension_matrix:
        assert re.search(r'UC-E\d{2}', cols) and re.search(r'P\d{2}', cols), rid
    return {'document_id': 'SRS-FAP-01', 'version': '3.1',
            'requirements': dict(counts), 'total_requirements': len(declared),
            'use_cases': len(uses), 'planned_test_cases': len(tests),
            'matrix_rows': len(matrix), 'extension_requirements':len(extensions),
            'total_requirement_ids':len(declared)+len(extensions),
            'total_planned_tests':len(tests)+len(extension_tests), 'validation': 'passed',
            'application_tests': 'not_run: specification-only task'}


def inline(paragraph, text: str):
    # Render bold spans and hyperlinks; retain readable technical tokens.
    tokens = re.split(r'(\*\*.*?\*\*|\[[^\]]+\]\([^)]+\))', text)
    for token in tokens:
        link = re.fullmatch(r'\[([^\]]+)\]\(([^)]+)\)', token)
        if link:
            label, target = link.groups()
            if not target.startswith('https://'):
                target = (SOURCE.parent / target).resolve().as_uri()
            rel = paragraph.part.relate_to(
                target,
                'http://schemas.openxmlformats.org/officeDocument/2006/relationships/hyperlink',
                is_external=True)
            hyp = OxmlElement('w:hyperlink')
            hyp.set(qn('r:id'), rel)
            run = OxmlElement('w:r')
            props = OxmlElement('w:rPr')
            color = OxmlElement('w:color'); color.set(qn('w:val'), '17365D')
            props.append(color); run.append(props)
            txt = OxmlElement('w:t'); txt.text = label; run.append(txt)
            hyp.append(run); paragraph._p.append(hyp)
        elif token.startswith('**') and token.endswith('**'):
            paragraph.add_run(token[2:-2]).bold = True
        else:
            paragraph.add_run(token.replace('`', ''))


def cell_properties(cell, header=False):
    cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.TOP
    tcpr = cell._tc.get_or_add_tcPr()
    margins = OxmlElement('w:tcMar')
    for name, value in [('top', '80'), ('bottom', '80'), ('left', '100'), ('right', '100')]:
        element = OxmlElement(f'w:{name}')
        element.set(qn('w:w'), value); element.set(qn('w:type'), 'dxa')
        margins.append(element)
    tcpr.append(margins)
    borders = OxmlElement('w:tcBorders')
    for edge in ('top', 'left', 'bottom', 'right'):
        element = OxmlElement(f'w:{edge}')
        for key, value in [('val', 'single'), ('sz', '4'), ('color', 'D9D9D9')]:
            element.set(qn(f'w:{key}'), value)
        borders.append(element)
    tcpr.append(borders)
    if header:
        shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), 'E9EEF4'); tcpr.append(shade)


def table(doc, lines):
    rows = [[col.strip() for col in line.strip().strip('|').split('|')] for line in lines]
    rows = [row for row in rows if not all(re.fullmatch(r':?-+:?', col) for col in row)]
    ncols = len(rows[0])
    assert all(len(row) == ncols for row in rows), rows
    tab = doc.add_table(rows=0, cols=ncols)
    tab.alignment = WD_TABLE_ALIGNMENT.CENTER
    tab.autofit = False
    if ncols == 2:
        widths = [4.9, 12.1]
    elif ncols == 3:
        widths = [3.3, 6.3, 7.4]
    elif ncols == 4:
        widths = [2.2, 6.1, 4.7, 4.0]
    else:
        widths = [1.8, 4.9, 2.7, 3.8, 3.8]
    for column, width in zip(tab.columns, widths):
        column.width = Cm(width)
    for ri, data in enumerate(rows):
        row = tab.add_row()
        no_split = OxmlElement('w:cantSplit'); row._tr.get_or_add_trPr().append(no_split)
        if ri == 0:
            repeat = OxmlElement('w:tblHeader'); row._tr.get_or_add_trPr().append(repeat)
        for ci, text in enumerate(data):
            cell = row.cells[ci]; cell.width = Cm(widths[ci]); cell_properties(cell, ri == 0)
            para = cell.paragraphs[0]
            para.paragraph_format.space_after = Pt(0)
            para.paragraph_format.line_spacing = 1.08
            inline(para, text.replace(',', ', ') if ncols == 5 else text)
            for run in para.runs:
                run.font.size = Pt(10)
                if ri == 0:
                    run.bold = True
            # Table header should remain attached to its first data row.
            para.paragraph_format.keep_with_next = ri == 0
    doc.add_paragraph().paragraph_format.space_after = Pt(0)


def build(text: str):
    doc = Document()
    sec = doc.sections[0]
    sec.page_width = Cm(21); sec.page_height = Cm(29.7)
    sec.top_margin = Cm(2.1); sec.bottom_margin = Cm(2.0)
    sec.left_margin = Cm(2); sec.right_margin = Cm(2)
    sec.header_distance = Cm(.9); sec.footer_distance = Cm(.9)
    sec.different_first_page_header_footer = True
    for name in ('Normal', 'Title', 'Subtitle', 'Heading 1', 'Heading 2', 'Heading 3'):
        style = doc.styles[name]
        style.font.name = 'Arial'
        style.font.color.rgb = RGBColor(0, 0, 0)
        style.element.get_or_add_rPr().rFonts.set(qn('w:eastAsia'), 'Arial')
        for border in style.element.findall('.//' + qn('w:pBdr')):
            border.getparent().remove(border)
        style.font.underline = False
    normal = doc.styles['Normal']
    normal.font.size = Pt(10.5)
    normal.paragraph_format.line_spacing = 1.12
    normal.paragraph_format.space_after = Pt(6)
    normal.paragraph_format.widow_control = True
    for name, size in [('Heading 1', 15), ('Heading 2', 12), ('Heading 3', 11)]:
        style = doc.styles[name]; style.font.size = Pt(size); style.font.bold = True
        style.paragraph_format.space_before = Pt(12)
        style.paragraph_format.space_after = Pt(6)
        style.paragraph_format.keep_with_next = True
    doc.styles['Title'].font.size = Pt(27)
    doc.styles['Subtitle'].font.size = Pt(15)
    head = sec.header.paragraphs[0]
    head.add_run('FAP  |  ĐẶC TẢ YÊU CẦU PHẦN MỀM  |  v3.1').font.size = Pt(8)
    head.paragraph_format.space_after = Pt(0)
    foot = sec.footer.paragraphs[0]; foot.alignment = WD_ALIGN_PARAGRAPH.RIGHT
    foot.add_run('SRS-FAP-01  •  Trang ').font.size = Pt(8)
    field = OxmlElement('w:fldSimple'); field.set(qn('w:instr'), 'PAGE'); foot._p.append(field)
    title = doc.add_paragraph('ĐẶC TẢ YÊU CẦU\nPHẦN MỀM', 'Title')
    title.paragraph_format.space_before = Pt(76)
    title.paragraph_format.space_after = Pt(26)
    doc.add_paragraph('SRS-FAP-01', 'Subtitle')
    topic = re.search(r'^\*\*(Xây dựng.+)\*\*$', text, re.M).group(1)
    p = doc.add_paragraph(topic)
    p.paragraph_format.space_after = Pt(28)
    p.runs[0].font.size = Pt(16); p.runs[0].bold = True
    doc.add_paragraph('Phiên bản 3.1  •  05/10/2026')
    doc.add_paragraph('Dự thảo cơ sở yêu cầu để rà soát và triển khai')
    doc.add_paragraph('Tham chiếu ISO/IEC/IEEE 29148:2018 và hướng dẫn công khai NASA; chưa ghi nhận phê duyệt baseline.')
    doc.add_page_break()
    doc.add_heading('Mục lục', 1)
    doc.add_paragraph('Các tiêu đề dưới đây tương ứng với Navigation Pane của Word.')
    toc_map = json.loads((QA / 'toc-pages.json').read_text(encoding='utf-8')) if (QA / 'toc-pages.json').exists() else {}
    for heading in re.findall(r'^## (.+)$', text, re.M):
        suffix = f'\t{toc_map[heading]}' if heading in toc_map else ''
        p = doc.add_paragraph(heading + suffix)
        p.paragraph_format.tab_stops.add_tab_stop(Cm(16.8), WD_TAB_ALIGNMENT.RIGHT, WD_TAB_LEADER.DOTS)
        p.paragraph_format.space_after = Pt(7)
    doc.add_page_break()
    lines = text.splitlines()
    start = next(i for i, line in enumerate(lines) if line.startswith('## 0.'))
    i = start
    while i < len(lines):
        line = lines[i].strip()
        if not line:
            i += 1; continue
        if line.startswith('|'):
            end = i + 1
            while end < len(lines) and lines[end].strip().startswith('|'):
                end += 1
            table(doc, lines[i:end]); i = end; continue
        match = re.match(r'^(#{2,4}) (.+)$', line)
        if match:
            level = len(match.group(1)) - 1
            p = doc.add_heading(match.group(2), level)
        else:
            p = doc.add_paragraph(); inline(p, line)
            if re.match(r'^\*\*(?:FR|BR|DR|IR|NFR)\d{2}\b', line):
                p.paragraph_format.keep_together = True
        i += 1
    doc.core_properties.title = 'SRS-FAP-01 — Đặc tả yêu cầu phần mềm'
    doc.core_properties.subject = topic
    doc.core_properties.author = 'Sinh viên thực hiện đề tài / trợ lý Codex'
    doc.core_properties.version = '3.1'
    OUTPUT.parent.mkdir(parents=True, exist_ok=True)
    doc.save(OUTPUT)


if __name__ == '__main__':
    source_text = SOURCE.read_text(encoding='utf-8')
    result = validate(source_text)
    build(source_text)
    QA.mkdir(parents=True, exist_ok=True)
    (QA / 'validation.json').write_text(json.dumps(result, ensure_ascii=False, indent=2), encoding='utf-8')
    print(json.dumps({**result, 'source': str(SOURCE), 'output': str(OUTPUT)}, ensure_ascii=False))
