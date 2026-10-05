import json
"""Build the factual proposal report for discussion with the supervisor."""
from pathlib import Path
from docx import Document
from docx.shared import Cm, Pt, RGBColor
from docx.oxml import OxmlElement
from docx.oxml.ns import qn
from docx.enum.table import WD_TABLE_ALIGNMENT, WD_CELL_VERTICAL_ALIGNMENT

ROOT = Path(__file__).resolve().parents[2]
OUT = ROOT / 'docs/research-platform/reports'
OUT.mkdir(parents=True, exist_ok=True)
doc = Document()
section = doc.sections[0]
section.page_height, section.page_width = Cm(29.7), Cm(21)
section.top_margin = section.bottom_margin = Cm(1.9)
section.left_margin = section.right_margin = Cm(2)
for style_name in ('Normal', 'Title', 'Heading 1', 'Heading 2'):
    style = doc.styles[style_name]
    style.font.name = 'Arial'
    style.font.color.rgb = RGBColor(0, 0, 0)
    style.font.size = Pt(10.5 if style_name == 'Normal' else 15 if style_name.startswith('Heading') else 21)
    style.paragraph_format.space_after = Pt(7)
    style.paragraph_format.line_spacing = 1.12


def para(text, style=None):
    return doc.add_paragraph(text, style)


def table(headers, rows, widths):
    tbl = doc.add_table(rows=1, cols=len(headers))
    tbl.alignment = WD_TABLE_ALIGNMENT.CENTER
    tbl.autofit = False
    for col, width in zip(tbl.columns, widths):
        col.width = Cm(width)
    for cell, title in zip(tbl.rows[0].cells, headers):
        cell.text = title
    repeat = OxmlElement('w:tblHeader')
    tbl.rows[0]._tr.get_or_add_trPr().append(repeat)
    for row in rows:
        for cell, text in zip(tbl.add_row().cells, row):
            cell.text = text
    props = tbl._tbl.tblPr
    borders = OxmlElement('w:tblBorders')
    for edge in ('top', 'left', 'bottom', 'right', 'insideH', 'insideV'):
        el = OxmlElement('w:' + edge)
        el.set(qn('w:val'), 'single'); el.set(qn('w:sz'), '4'); el.set(qn('w:color'), 'D9D9D9')
        borders.append(el)
    props.append(borders)
    for index, row in enumerate(tbl.rows):
        for cell, width in zip(row.cells, widths):
            cell.width = Cm(width)
            cell.vertical_alignment = WD_CELL_VERTICAL_ALIGNMENT.CENTER
            tcPr = cell._tc.get_or_add_tcPr()
            shade = OxmlElement('w:shd'); shade.set(qn('w:fill'), '163A59' if index == 0 else 'F0F4F7' if index % 2 == 0 else 'FFFFFF'); tcPr.append(shade)
            margins = OxmlElement('w:tcMar')
            for edge in ('top', 'left', 'bottom', 'right'):
                el = OxmlElement('w:' + edge); el.set(qn('w:w'), '100'); el.set(qn('w:type'), 'dxa'); margins.append(el)
            tcPr.append(margins)
            for p in cell.paragraphs:
                p.paragraph_format.space_after = Pt(3)
                p.paragraph_format.line_spacing = 1.05
                for run in p.runs:
                    run.font.size = Pt(10)
                    if index == 0:
                        run.bold = True; run.font.color.rgb = RGBColor(255,255,255)
    para('')



for style_name in ('Title', 'Subtitle'):
    props=doc.styles[style_name]._element.find(qn('w:pPr'))
    if props is not None:
        for border in list(props.findall(qn('w:pBdr'))): props.remove(border)

data=json.loads((OUT/'noi-dung-de-xuat.json').read_text(encoding='utf-8'))
doc.core_properties.title = data['topic']
doc.core_properties.subject = 'Phương pháp dữ liệu và kế hoạch project BCTC–cổ tức'
for i,page in enumerate(data['pages']):
    if i: doc.add_page_break()
    else:
        para(data['title'], 'Title')
        para('Báo cáo định hướng ngày '+data['date'])
        para(data['topic'])
    para(page['heading'], 'Heading 1')
    for block in page['blocks']:
        if 'h' in block: para(block['h'], 'Heading 2')
        elif 'p' in block: para(block['p'])
        else:
            t=block['table'];table(t['headers'],t['rows'],t['widths'])
doc.save(OUT/'bao-cao-de-xuat.docx')
print('Saved proposal DOCX from shared content')
